import streamlit as st
from fpdf import FPDF
from docx import Document
import base64
from PIL import Image
from io import BytesIO
import os
import shutil

# ✅ Page Configuration
st.set_page_config(page_title="GenAI Resume Builder", layout="centered")

# 🎨 Custom CSS Styling
st.markdown("""
    <style>
        @keyframes fadeIn {
            from {opacity: 0;}
            to {opacity: 1;}
        }

        .stApp {
            background-color: #000;
            color: #ffffff;
            font-family: 'Segoe UI', sans-serif;
            animation: fadeIn 1s ease-in-out;
            overflow: hidden;
        }

        .main > div {
            background-color: #1a1a1a;
            padding: 30px;
            border-radius: 12px;
            box-shadow: 0 4px 20px rgba(255,0,0,0.3);
            color: #ffffff;
            animation: fadeIn 1.5s ease-in-out;
        }

        input, textarea {
            background-color: #333 !important;
            color: #fff !important;
            border: 1px solid #ffffff !important;
            border-radius: 10px !important;
            padding: 10px !important;
            transition: transform 0.2s ease-in-out, box-shadow 0.2s ease-in-out;
        }

        input:hover, textarea:hover {
            border-color: #ffffff !important;
            transform: scale(1.03);
            box-shadow: 0 0 10px rgba(255,255,255,0.2);
        }

        .preview:hover {
            transform: scale(1.01);
            box-shadow: 0 0 15px rgba(255,0,0,0.5);
            transition: 0.3s ease-in-out;
        }
    </style>
""", unsafe_allow_html=True)

# 🧬 Title and Subtitle
st.title("GenAI Resume Builder")
st.markdown("<h4 style='color:#ffffff;'>Craft your resume in style. Fill in your details below and download your custom resume!</h4>", unsafe_allow_html=True)

# 📷 Upload Image
st.header("📷 Upload Profile Picture (Required)")
image_file = st.file_uploader("Upload your photo (JPG or PNG)", type=['jpg', 'jpeg', 'png'])

# 📅 Personal Information
st.header("🧑 Personal Information")
name = st.text_input("Full Name", placeholder="e.g. John Doe")
job_title = st.text_input("Job Title You're Applying For", placeholder="e.g. Software Engineer")

# 💼 Experience Section
st.header("🏢 Work Experience")
experience = st.text_area("List your work experience (bullet points, one per line)", placeholder="- Company A: Role\n- Company B: Role")

# 🎓 Education Section
st.header("🏫 Education")
school = st.text_input("School Name", placeholder="e.g. ABC High School")
college = st.text_input("College Name", placeholder="e.g. XYZ University")
degree = st.text_input("Degree Obtained", placeholder="e.g. B.Tech in Computer Science")
edu_summary = st.text_area("Education Summary (bullet points, one per line)", placeholder="- GPA: 8.5\n- Dean’s List for 3 years")

# 🧠 Skills Section
st.header("🛠️ Skills")
skills = st.text_area("List your skills (comma-separated)", placeholder="Python, HTML, CSS, JavaScript")

# 📱 Contact Info
st.header("📬 Contact Information")
linkedin = st.text_input("LinkedIn URL", placeholder="[https://linkedin.com/in/yourname](https://linkedin.com/in/yourname)")
email = st.text_input("Email Address", placeholder="[example@email.com](mailto:example@email.com)")
facebook = st.text_input("Facebook Profile URL", placeholder="[https://facebook.com/yourprofile](https://facebook.com/yourprofile)")
manual_contact = st.text_area("Other Contact Info (e.g., phone, address)", placeholder="Phone: +91-1234567890\nAddress: Your city, country")

# 🚀 Generate Resume
if st.button("⚙️ Generate Resume"):
    if not all([name, job_title, experience, skills, email, degree]):
        st.warning("❗ Please fill in all required fields before generating the resume.")
    elif not image_file:
        st.warning("❗ Profile picture is required!")
    else:
        exp_lines = experience.strip().split('\n')
        edu_lines = edu_summary.strip().split('\n')
        skills_list = [s.strip() for s in skills.split(',')]

        # Process Profile Image
        img = Image.open(image_file).convert("RGB")
        img = img.resize((120, 120))
        buffered = BytesIO()
        img.save(buffered, format="PNG")
        img_base64 = base64.b64encode(buffered.getvalue()).decode()
        img_html = f'<img src="data:image/png;base64,{img_base64}" style="border-radius: 50%;width:120px;height:120px;object-fit:cover;"/>' 

        # Resume Preview HTML
        resume_html = f"""
        <div class='preview' style="background:#fff;color:#000;padding:30px;border-radius:10px;">
            <h1 style="text-align:center;color:#d10000;">Resume</h1>
            <div style="text-align:center;">{img_html}</div>
            <h2 style="color:#d10000;">{name}</h2>
            <h4>{job_title}</h4>
            <hr>
            <h4 style="color:#d10000;">Summary</h4>
            <p>A highly motivated professional applying for the role of {job_title}.</p>

            <h4 style="color:#d10000;">Experience</h4>
            <ul>{''.join([f'<li>{line}</li>' for line in exp_lines])}</ul>

            <h4 style="color:#d10000;">Education</h4>
            <p><b>School:</b> {school}<br><b>College:</b> {college}<br><b>Degree:</b> {degree}</p>
            <ul>{''.join([f'<li>{line}</li>' for line in edu_lines])}</ul>

            <h4 style="color:#d10000;">Skills</h4>
            <p>{' | '.join(skills_list)}</p>

            <h4 style="color:#d10000;">Contact</h4>
            <p><b>Email:</b> {email}<br><b>LinkedIn:</b> {linkedin}<br><b>Facebook:</b> {facebook}<br><b>Other:</b> {manual_contact}</p>
        </div>
        """
        st.markdown("### 🧾 Resume Preview")
        st.components.v1.html(resume_html, height=800, scrolling=True)

        # ✅ PDF Generation
        pdf = FPDF()
        pdf.add_page()

        img_path = "profile_pic_temp.png"
        img.save(img_path, format="PNG")
        pdf.image(img_path, x=10, y=8, w=30)

        pdf.set_xy(45, 10)
        pdf.set_font("Arial", 'B', 16)
        pdf.cell(0, 10, name, ln=True)
        pdf.set_font("Arial", '', 12)
        pdf.cell(0, 10, job_title, ln=True)
        pdf.ln(10)

        pdf.set_font("Arial", 'B', 14)
        pdf.cell(0, 10, "Summary", ln=True)
        pdf.set_font("Arial", '', 12)
        pdf.multi_cell(0, 10, f"A highly motivated professional applying for the role of {job_title}.")

        pdf.set_font("Arial", 'B', 14)
        pdf.cell(0, 10, "Experience", ln=True)
        pdf.set_font("Arial", '', 12)
        for exp in exp_lines:
            pdf.multi_cell(0, 10, f"- {exp.strip()}")

        pdf.set_font("Arial", 'B', 14)
        pdf.cell(0, 10, "Education", ln=True)
        pdf.set_font("Arial", '', 12)
        pdf.multi_cell(0, 10, f"School: {school}\nCollege: {college}\nDegree: {degree}")
        for edu in edu_lines:
            pdf.multi_cell(0, 10, f"- {edu.strip()}")

        pdf.set_font("Arial", 'B', 14)
        pdf.cell(0, 10, "Skills", ln=True)
        pdf.set_font("Arial", '', 12)
        pdf.multi_cell(0, 10, ', '.join(skills_list))

        pdf.set_font("Arial", 'B', 14)
        pdf.cell(0, 10, "Contact", ln=True)
        pdf.set_font("Arial", '', 12)
        pdf.multi_cell(0, 10, f"Email: {email}\nLinkedIn: {linkedin}\nFacebook: {facebook}\nOther: {manual_contact}")

        # PDF generation complete
        pdf_bytes = pdf.output(dest='S').encode('latin1', 'replace')

        # Allow PDF download
        st.download_button(
            label="📥 Download Resume (PDF)",
            data=pdf_bytes,
            file_name=f"{name.replace(' ', '_')}_resume.pdf",
            mime="application/pdf"
        )

        # Clean up temp file
        if os.path.exists(img_path):
            try:
                os.remove(img_path)
                print(f"Temporary file {img_path} has been deleted.")
            except Exception as e:
                print(f"Error deleting temporary file {img_path}: {e}")
